"""Check the Excel and CSV files have the promised formatting."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

from price_tracker.exporter import GBP_FORMAT, write_csv, write_excel
from price_tracker.models import Change, ChangeType
from tests.conftest import make_book

BOOKS = [make_book("A Very Long Book Title For Width Testing", 12.5), make_book("B", 3.0, in_stock=False)]
CHANGES = [Change(ChangeType.PRICE_UP, "B", "Mystery", BOOKS[1].url, 2.0, 3.0)]


def test_excel_formatting(tmp_path: Path) -> None:
    path = write_excel(BOOKS, CHANGES, tmp_path / "out.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Books", "Changes"]

    books = workbook["Books"]
    assert [c.value for c in books[1]][:3] == ["Title", "Price", "Rating"]
    assert all(cell.font.bold for cell in books[1])
    assert books.freeze_panes == "A2"
    assert books["B2"].value == 12.5 and books["B2"].number_format == GBP_FORMAT
    assert books["D3"].value == "Out of stock"
    assert books.column_dimensions["A"].width > books.column_dimensions["C"].width

    changes = workbook["Changes"]
    assert changes.freeze_panes == "A2"
    assert changes["A2"].value == "Price increase"
    assert changes["F2"].value == 1.0 and changes["F2"].number_format == GBP_FORMAT
    assert changes["G2"].value == 0.5


def test_excel_baseline_note(tmp_path: Path) -> None:
    path = write_excel(BOOKS, [], tmp_path / "out.xlsx", baseline=True)
    assert "Baseline" in load_workbook(path)["Changes"]["A2"].value


def test_csv_export(tmp_path: Path) -> None:
    path = write_csv(BOOKS, tmp_path / "out.csv")
    with path.open(encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 2
    assert rows[0]["price_gbp"] == "12.50"
    assert rows[1]["stock_status"] == "Out of stock"
