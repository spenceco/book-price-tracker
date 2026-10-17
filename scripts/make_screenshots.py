"""Create portfolio PNGs: one per Excel sheet, plus a terminal run.

Needs LibreOffice (``soffice``), ``pdftoppm`` (poppler-utils) and Pillow.

    python scripts/make_screenshots.py output/books_....xlsx run.log docs/screenshots

How it works:
* Excel -> for each sheet, save a trimmed copy (first rows only, with grid
  lines and row/column headings switched on), convert it to PDF with
  LibreOffice headless, rasterise page 1 with pdftoppm, then crop and frame it.
* Terminal -> the captured console output (including colour codes) is drawn
  onto a dark "terminal window" with Pillow.
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageChops, ImageDraw, ImageFont

MONO = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"
MONO_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf"
SANS = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"

# --- Excel sheets -------------------------------------------------------


def sheet_to_png(xlsx: Path, sheet: str, rows: int, out: Path, workdir: Path) -> Path:
    """Render the top ``rows`` data rows of one sheet to a cropped PNG."""
    workbook = load_workbook(xlsx)
    for name in workbook.sheetnames:
        if name != sheet:
            del workbook[name]
    ws = workbook[sheet]
    if ws.max_row > rows + 1:
        ws.delete_rows(rows + 2, ws.max_row)
    ws.print_options.gridLines = True
    ws.print_options.headings = True  # show A, B, C... and 1, 2, 3...
    ws.page_setup.fitToHeight = 1
    ws.print_title_rows = None
    ws.print_area = ws.dimensions
    ws.oddHeader.center.text = ws.oddFooter.center.text = ""
    for side in ("left", "right", "top", "bottom"):
        setattr(ws.page_margins, side, 0.2)
    trimmed = workdir / f"{sheet}.xlsx"
    workbook.save(trimmed)

    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(workdir), str(trimmed)],
        check=True, capture_output=True, timeout=180,
    )
    stem = workdir / f"{sheet}-page"
    subprocess.run(["pdftoppm", "-png", "-r", "220", "-f", "1", "-l", "1", "-singlefile",
                    str(trimmed.with_suffix(".pdf")), str(stem)], check=True)
    image = _crop_whitespace(Image.open(f"{stem}.png").convert("RGB"))
    framed = _window_frame(image, f"{xlsx.name}  —  {sheet} sheet", dark=False)
    framed.save(out, optimize=True)
    return out


def _crop_whitespace(image: Image.Image, pad: int = 12) -> Image.Image:
    background = Image.new("RGB", image.size, "white")
    box = ImageChops.difference(image, background).getbbox()
    if not box:
        return image
    left, top, right, bottom = box
    return image.crop((max(left - pad, 0), max(top - pad, 0),
                       min(right + pad, image.width), min(bottom + pad, image.height)))


def _window_frame(content: Image.Image, title: str, dark: bool) -> Image.Image:
    """Wrap an image in a simple desktop-window frame with a title bar."""
    bar, margin = 44, 28
    chrome = "#2b2b2b" if dark else "#e9ecef"
    canvas = Image.new("RGB", (content.width + 2 * margin, content.height + bar + 2 * margin), "#f4f6f8")
    draw = ImageDraw.Draw(canvas)
    draw.rounded_rectangle((margin - 1, margin - 1, margin + content.width, margin + bar + content.height),
                           radius=10, fill=chrome, outline="#c7cbd1")
    for i, colour in enumerate(("#ff5f57", "#febc2e", "#28c840")):
        x = margin + 20 + i * 22
        draw.ellipse((x, margin + 16, x + 12, margin + 28), fill=colour)
    font = ImageFont.truetype(SANS, 17)
    draw.text((margin + content.width / 2, margin + bar / 2), title, anchor="mm",
              font=font, fill="#d0d0d0" if dark else "#40454d")
    canvas.paste(content, (margin, margin + bar))
    return canvas


# --- Terminal -----------------------------------------------------------

ANSI_RE = re.compile(r"\x1b\[([0-9;]*)m")
PALETTE = {"31": "#ff6b6b", "32": "#69db7c", "33": "#ffd43b", "36": "#66d9e8"}
FG, BG = "#d8dee9", "#1e222a"


def terminal_to_png(log_text: str, command: str, out: Path) -> Path:
    """Draw console output (with ANSI colours) as a terminal screenshot."""
    font, bold = ImageFont.truetype(MONO, 20), ImageFont.truetype(MONO_BOLD, 20)
    char_w = font.getlength("M")
    line_h = 30
    lines = [f"\x1b[32m$\x1b[0m {command}"] + log_text.rstrip("\n").splitlines()
    width = int(max(len(ANSI_RE.sub("", line)) for line in lines) * char_w) + 48
    image = Image.new("RGB", (max(width, 900), len(lines) * line_h + 36), BG)
    draw = ImageDraw.Draw(image)

    for row, line in enumerate(lines):
        x, y = 24.0, 18 + row * line_h
        colour, is_bold, dim = FG, False, False
        # "SUMMARY" lines get a highlight so the headline stands out
        if "SUMMARY:" in line:
            is_bold = True
        for index, part in enumerate(ANSI_RE.split(line)):
            if index % 2:  # an SGR code like "36" or "0"
                for code in (part or "0").split(";"):
                    if code == "0":
                        colour, dim = FG, False
                        is_bold = "SUMMARY:" in line
                    elif code == "1":
                        is_bold = True
                    elif code == "2":
                        dim = True
                    elif code in PALETTE:
                        colour = PALETTE[code]
                continue
            fill = "#7a8291" if dim else colour
            draw.text((x, y), part, font=bold if is_bold else font, fill=fill)
            x += font.getlength(part)
    _window_frame(image, "price-tracker — bash", dark=True).save(out, optimize=True)
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate portfolio screenshots.")
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("terminal_log", type=Path, help="captured console output")
    parser.add_argument("out_dir", type=Path)
    parser.add_argument("--command", default="python -m price_tracker")
    parser.add_argument("--rows", type=int, default=25, help="data rows shown from the Books sheet")
    args = parser.parse_args()
    if not shutil.which("soffice"):
        raise SystemExit("LibreOffice (soffice) is required")
    args.out_dir.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        print(sheet_to_png(args.xlsx, "Books", args.rows, args.out_dir / "excel_books.png", work))
        print(sheet_to_png(args.xlsx, "Changes", 40, args.out_dir / "excel_changes.png", work))
    log_text = args.terminal_log.read_text(encoding="utf-8")
    print(terminal_to_png(log_text, args.command, args.out_dir / "terminal_run.png"))


if __name__ == "__main__":
    main()
